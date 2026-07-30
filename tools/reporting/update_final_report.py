from __future__ import annotations

import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
REPORT = (
    ROOT
    / "deliverables"
    / "informe_final"
    / "Parte_Cono_de_Carga_Informe_Final.docx"
)
TRAINING_FIGURE = (
    ROOT
    / "deliverables"
    / "informe_final"
    / "figuras"
    / "Figura_C4_entrenamiento_mlflow.png"
)


PARAGRAPH_REPLACEMENTS = {
    "La implementación integra Python, OpenCV": (
        "La implementación integra Python, OpenCV, NumPy, SciPy, Pandas, "
        "Ultralytics, Docker, NVIDIA CUDA, MLflow y Streamlit. El repositorio "
        "de percepción registra un entrenamiento reproducible con 240 imágenes "
        "y conserva su best.pt como checkpoint alternativo. Para la operación "
        "se seleccionó el modelo principal PSINet. En una prueba técnica "
        "balanceada de 90 imágenes, el pipeline obtiene 88,89 % de exactitud "
        "global: 73,33 % para centrado, 93,33 % para izquierda y 100 % para "
        "derecha. Las medianas de pico, simetría y centroide son cercanas a "
        "cero para centrado, negativas para izquierda y positivas para derecha, "
        "lo que aporta una explicación geométrica consistente. Una muestra "
        "descriptiva de 1.248 imágenes contiene 1.183 casos evaluables y 65 no "
        "evaluables. Los resultados confirman que la segmentación condiciona la "
        "clasificación. Las salidas representan asimetría visual interna, no "
        "balance físico, distribución real del material ni desplazamiento "
        "respecto de la tolva."
    ),
    "The implementation combines Python, OpenCV": (
        "The implementation combines Python, OpenCV, NumPy, SciPy, Pandas, "
        "Ultralytics, Docker, NVIDIA CUDA, MLflow, and Streamlit. The repository "
        "records a reproducible training run with 240 images and retains its "
        "best.pt as an alternative checkpoint, while PSINet is used as the "
        "primary operational segmenter. In a balanced technical evaluation with "
        "90 visually selected images, the complete pipeline reaches 88.89% "
        "overall accuracy. Median apex, area-symmetry, and centroid indicators "
        "remain close to zero for centered loads, negative for left cases, and "
        "positive for right cases. A separate descriptive sample of 1,248 images "
        "contains 1,183 evaluable and 65 non-evaluable cases."
    ),
    "5.3 Entrenamiento reproducible": "5.3 Entrenamiento reproducible del checkpoint alternativo",
    "El entrenamiento parte desde la definición yolov8n-seg.yaml": (
        "El entrenamiento reproducible parte desde yolov8n-seg.yaml, sin pesos "
        "preentrenados. El uso de pretrained=false demuestra que el repositorio "
        "entrena un segmentador y no se limita a recibir pesos terminados. "
        "MLflow registra la configuración, el conjunto, los parámetros, las "
        "métricas y los artefactos best.pt y last.pt en el servidor definido "
        "mediante .env.mlflow [11]. Tras la evaluación extremo a extremo, "
        "best.pt se conserva como alternativa y PSINet queda como modelo "
        "principal de segmentación."
    ),
    "Tabla C.5. Configuración del entrenamiento.": (
        "Tabla C.5. Configuración del entrenamiento del checkpoint alternativo."
    ),
    "Figura C.4. Seguimiento del entrenamiento": (
        "Figura C.4. Seguimiento y trazabilidad del entrenamiento del checkpoint alternativo."
    ),
    "5.4 Resultados del modelo de percepción": "5.4 Modelo de percepción operativo y referencia de entrenamiento",
    "El mejor punto del entrenamiento alcanza un mAP": (
        "En el entrenamiento del checkpoint alternativo, el mejor punto alcanza "
        "un mAP de máscara entre IoU 0,50 y 0,95 de 0,94142 en la época 87. Al "
        "finalizar se registra mAP50 de máscara igual a 0,995 y mAP50-95 igual "
        "a 0,91647. Estas métricas muestran buen ajuste sobre la validación "
        "disponible, pero las 36 imágenes y su posible relación temporal impiden "
        "interpretarlas como una estimación definitiva de generalización. La "
        "selección operativa se basó además en el desempeño del pipeline "
        "completo, donde PSINet obtuvo los mejores resultados."
    ),
    "El modelo fue entrenado para identificar la carga visible": (
        "El pipeline utiliza como segmentador principal "
        "perception/models/primary/psinet_carga_visible.pt y selecciona su "
        "clase material (class_id=1). El checkpoint "
        "perception/models/alternative/best.pt se mantiene para respaldo y "
        "comparaciones controladas con su clase Carga (class_id=0). Durante la "
        "inferencia se aplica una confianza de 0,35 y el modo largest conserva "
        "la instancia de carga de mayor área. La predicción se rasteriza como "
        "una única máscara RGB, contrato de entrada del analítico geométrico."
    ),
    "Los umbrales se mantienen en el archivo de configuración": (
        "Los umbrales se mantienen en el archivo de configuración para "
        "conservar trazabilidad y permitir recalibración ante nuevas cámaras o "
        "condiciones. El analítico interpreta la máscara entregada por el modelo "
        "principal de percepción y no intenta reconstruir regiones omitidas ni "
        "corregir contornos erróneos."
    ),
    "El modelo de percepción entrenado se evalúa junto": (
        "El modelo principal PSINet se evalúa junto con el analítico sobre las "
        "90 imágenes de la prueba técnica. Todas resultan evaluables y 80 se "
        "clasifican correctamente, equivalentes a 88,89 % de exactitud global. "
        "Este valor representa el mejor resultado extremo a extremo obtenido y "
        "corresponde a la configuración utilizada por el pipeline."
    ),
    "Fuente: elaboración propia a partir de la prueba técnica": (
        "Fuente: elaboración propia a partir de la prueba técnica balanceada de "
        "90 imágenes procesada con el modelo principal de percepción."
    ),
    "Tabla C.9. Métricas por clase del modelo de percepción.": (
        "Tabla C.9. Métricas de clasificación por clase."
    ),
    "La prueba técnica confirma que la combinación": (
        "La prueba técnica confirma que la combinación de segmentación y "
        "geometría puede discriminar tendencias visuales laterales en el "
        "conjunto evaluado. El modelo principal PSINet, integrado como "
        "segmentador operativo, alcanza 88,89 % de exactitud global. La revisión "
        "de errores muestra que la percepción sigue siendo un factor limitante, "
        "pues las reglas geométricas no recuperan regiones omitidas ni remueven "
        "regiones incorrectas sin información adicional."
    ),
    "El repositorio mantiene configuración y código": (
        "El repositorio mantiene configuración y código separados de datos y "
        "resultados. Los umbrales del analítico residen en analytic_config.yaml; "
        "los parámetros de percepción, en YAML; y las credenciales de MLflow, "
        "en un archivo de entorno excluido de Git. El modelo principal PSINet y "
        "el checkpoint alternativo best.pt se conservan en rutas explícitas, "
        "mientras MLflow mantiene la trazabilidad del entrenamiento."
    ),
    "La evaluación técnica obtiene 88,89 %": (
        "La evaluación técnica obtiene 88,89 % de exactitud global con el modelo "
        "principal PSINet. El análisis de errores confirma que la calidad y "
        "cobertura de la máscara condicionan el resultado extremo a extremo. Los "
        "65 descartes en la muestra de 1.248 imágenes también muestran la "
        "utilidad de declarar no evaluables los casos sin evidencia suficiente."
    ),
    "Versionar explícitamente el modelo de percepción": (
        "Versionar explícitamente el modelo principal y el checkpoint "
        "alternativo, y registrar junto a cada corrida la ruta del peso, la "
        "clase seleccionada, el umbral de confianza y el conjunto de evaluación."
    ),
}


def replace_paragraph_text(paragraph, replacement: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)


def update_text(document: Document) -> None:
    found = set()
    for prefix, replacement in PARAGRAPH_REPLACEMENTS.items():
        for paragraph in document.paragraphs:
            if paragraph.text.startswith(prefix):
                replace_paragraph_text(paragraph, replacement)
                found.add(prefix)
                break
            if paragraph.text == replacement:
                found.add(prefix)
                break

    missing = set(PARAGRAPH_REPLACEMENTS) - found
    if missing:
        raise RuntimeError(f"No se encontraron párrafos esperados: {sorted(missing)}")

    alternatives_table = document.tables[0]
    alternatives_table.cell(4, 3).text = "Solución implementada."

    training_table = document.tables[4]
    training_table.cell(1, 1).text = "yolov8n-seg (alternativo)"
    training_table.cell(1, 2).text = "Checkpoint reproducible conservado como respaldo."


def replace_embedded_figure(docx_path: Path, figure_path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="cone_report_") as temp_dir:
        temp = Path(temp_dir)
        unpacked = temp / "unpacked"
        unpacked.mkdir()
        with zipfile.ZipFile(docx_path) as archive:
            archive.extractall(unpacked)

        embedded = unpacked / "word" / "media" / "image4.png"
        if not embedded.exists():
            raise FileNotFoundError(f"No se encontró la figura C4 embebida: {embedded}")
        shutil.copy2(figure_path, embedded)

        rebuilt = temp / docx_path.name
        with zipfile.ZipFile(rebuilt, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(unpacked.rglob("*")):
                if path.is_file():
                    archive.write(path, path.relative_to(unpacked).as_posix())
        shutil.copy2(rebuilt, docx_path)


def main() -> None:
    if not REPORT.exists():
        raise FileNotFoundError(REPORT)
    if not TRAINING_FIGURE.exists():
        raise FileNotFoundError(TRAINING_FIGURE)

    document = Document(REPORT)
    update_text(document)
    document.save(REPORT)
    replace_embedded_figure(REPORT, TRAINING_FIGURE)
    print(f"Informe actualizado: {REPORT}")


if __name__ == "__main__":
    main()
