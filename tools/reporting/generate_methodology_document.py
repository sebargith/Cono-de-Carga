from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = (
    ROOT
    / "deliverables"
    / "metodologia"
    / "Metodologia_general_cono_de_carga.docx"
)
FIGURES = ROOT / "deliverables" / "informe_final" / "figuras"

BLUE = RGBColor(31, 77, 120)
ACCENT = RGBColor(46, 116, 181)
INK = RGBColor(24, 35, 48)
MUTED = RGBColor(91, 101, 112)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
LIGHT_GREEN = "EAF4EE"
WHITE = RGBColor(255, 255, 255)
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(
    cell,
    *,
    top: int = 80,
    start: int = 120,
    bottom: int = 80,
    end: int = 120,
) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError("Los anchos de tabla deben sumar 9360 DXA.")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    properties = table._tbl.tblPr

    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    width.set(qn("w:type"), "dxa")

    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(widths_dxa[index]))
            cell_width.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_header_row(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def format_table(table, widths_dxa: list[int]) -> None:
    set_table_geometry(table, widths_dxa)
    repeat_header_row(table.rows[0])
    for column, cell in enumerate(table.rows[0].cells):
        set_cell_shading(cell, LIGHT_BLUE)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_run_font(run, size=9.5, color=INK, bold=True)
    for row in table.rows[1:]:
        for column, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(run, size=9.2, color=INK)
            if column == 0:
                set_cell_shading(cell, LIGHT_GRAY)


def set_repeatable_list_style(document: Document, style_name: str) -> None:
    style = document.styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.5)
    style.paragraph_format.first_line_indent = Inches(-0.25)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.167


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Title", 28, BLUE, 0, 8),
        ("Subtitle", 14, MUTED, 0, 16),
        ("Heading 1", 16, ACCENT, 16, 8),
        ("Heading 2", 13, ACCENT, 12, 6),
        ("Heading 3", 12, BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if style_name == "Heading 1":
            style.paragraph_format.page_break_before = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.color.rgb = MUTED
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False

    set_repeatable_list_style(document, "List Bullet")
    set_repeatable_list_style(document, "List Number")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.append(begin)
    field_run.append(instruction)
    field_run.append(separate)
    field_run.append(result)
    field_run.append(end)
    paragraph._p.append(field_run)


def configure_sections(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("METODOLOGÍA GENERAL  |  CONO DE CARGA")
        set_run_font(run, size=8.5, color=MUTED, bold=True)

        footer = section.footer
        footer.is_linked_to_previous = False
        add_page_number(footer.paragraphs[0])


def add_body_paragraph(
    document: Document,
    text: str,
    *,
    bold_lead: str | None = None,
    keep_with_next: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, color=INK)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run, color=INK)


def add_callout(
    document: Document,
    title: str,
    text: str,
    *,
    fill: str = LIGHT_GREEN,
) -> None:
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    repeat_header_row(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.paragraphs[0].paragraph_format.space_after = Pt(3)
    title_run = cell.paragraphs[0].add_run(title)
    set_run_font(title_run, size=10.5, color=BLUE, bold=True)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(
    document: Document,
    filename: str,
    caption: str,
    alt_text: str,
    *,
    width: float = 6.25,
) -> None:
    path = FIGURES / filename
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    inline = document.inline_shapes[-1]._inline
    inline.docPr.set("descr", alt_text)
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, color=MUTED, italic=True)


def math_text(text: str) -> OxmlElement:
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    return run


def math_subscript(base: str, subscript: str) -> OxmlElement:
    node = OxmlElement("m:sSub")
    base_node = OxmlElement("m:e")
    base_node.append(math_text(base))
    subscript_node = OxmlElement("m:sub")
    subscript_node.append(math_text(subscript))
    node.append(base_node)
    node.append(subscript_node)
    return node


def append_math_parts(parent: OxmlElement, parts: list[str | tuple[str, str]]) -> None:
    for part in parts:
        if isinstance(part, tuple):
            parent.append(math_subscript(part[0], part[1]))
        else:
            parent.append(math_text(part))


def math_fraction(
    numerator: list[str | tuple[str, str]],
    denominator: list[str | tuple[str, str]],
) -> OxmlElement:
    fraction = OxmlElement("m:f")
    numerator_node = OxmlElement("m:num")
    append_math_parts(numerator_node, numerator)
    denominator_node = OxmlElement("m:den")
    append_math_parts(denominator_node, denominator)
    fraction.append(numerator_node)
    fraction.append(denominator_node)
    return fraction


def add_fraction_equation(
    document: Document,
    left: list[str | tuple[str, str]],
    numerator: list[str | tuple[str, str]],
    denominator: list[str | tuple[str, str]],
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    math_paragraph = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    append_math_parts(math, left)
    math.append(math_text(" = "))
    math.append(math_fraction(numerator, denominator))
    math_paragraph.append(math)
    paragraph._p.append(math_paragraph)


def add_cover(document: Document) -> None:
    document.add_paragraph().paragraph_format.space_after = Pt(20)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("PROYECTO MEMORIA")
    set_run_font(run, size=11, color=ACCENT, bold=True)
    kicker.paragraph_format.space_after = Pt(28)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Metodología general del sistema de Cono de Carga")
    set_run_font(title_run, size=28, color=BLUE, bold=True)

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        "Segmentación de la carga visible, analítica geométrica y agregación "
        "temporal en camiones CAEX"
    )
    set_run_font(subtitle_run, size=14, color=MUTED)

    document.add_paragraph().paragraph_format.space_after = Pt(18)

    metadata = document.add_table(rows=4, cols=2)
    metadata.style = "Table Grid"
    labels = (
        ("Proyecto", "Memoria: Cono de Carga y Granulometría"),
        ("Instituciones", "Universidad de Concepción y PSINet"),
        ("Origen operacional de datos", "Registros audiovisuales de Collahuasi"),
        ("Versión", "Metodología general, julio de 2026"),
    )
    for row, (label, value) in zip(metadata.rows, labels):
        row.cells[0].text = label
        row.cells[1].text = value
    format_table(metadata, [2500, 6860])

    document.add_paragraph().paragraph_format.space_after = Pt(18)
    add_callout(
        document,
        "Propósito del documento",
        "Describir el método completo mediante el cual imágenes o videos "
        "posteriores de camiones CAEX se transforman en máscaras de carga, "
        "indicadores geométricos internos, clases visuales por frame y un "
        "veredicto agregado por video.",
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(36)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(
        "La metodología evalúa la forma visible de la carga. No estima masa, "
        "balance físico ni distribución real del payload."
    )
    set_run_font(run, size=10, color=MUTED, italic=True)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_scope(document: Document) -> None:
    document.add_heading("1. Objetivo y alcance metodológico", level=1)
    add_body_paragraph(
        document,
        "El sistema tiene por objetivo describir y clasificar la forma visible "
        "de la carga transportada por camiones CAEX observados desde atrás. La "
        "clasificación se expresa como centrado, izquierda, derecha o no "
        "evaluable y se construye mediante referencias internas a la propia "
        "máscara segmentada."
    )
    add_body_paragraph(
        document,
        "El método se diseñó en dos niveles. Primero, cada frame se procesa de "
        "forma independiente para obtener descriptores geométricos. Segundo, "
        "los resultados válidos de una misma secuencia se agregan para producir "
        "un veredicto temporal por video."
    )

    document.add_heading("1.1 Principios de diseño", level=2)
    for text in (
        "Separar percepción y analítica: la segmentación produce una máscara y "
        "el analítico interpreta únicamente esa geometría.",
        "Usar medidas adimensionales relativas al bounding box de la carga para "
        "reducir la dependencia del tamaño aparente del camión.",
        "Declarar no evaluables los frames con evidencia insuficiente, máscaras "
        "vacías o cargas truncadas por los bordes de la imagen.",
        "Mantener trazabilidad mediante configuraciones YAML, manifests, CSV, "
        "JSON y overlays visuales.",
        "Evitar interpretaciones físicas no sustentadas por la imagen "
        "bidimensional ni por una referencia explícita de la tolva.",
    ):
        add_bullet(document, text)

    document.add_heading("1.2 Arquitectura general", level=2)
    add_body_paragraph(
        document,
        "La solución se divide en extracción de frames, segmentación, adaptación "
        "de máscaras, analítica geométrica, agregación por video y presentación "
        "de resultados. Los contratos entre etapas permiten reemplazar o "
        "validar un componente sin modificar los demás."
    )
    add_figure(
        document,
        "Figura_C2_arquitectura.jpg",
        "Figura 1. Arquitectura general de la metodología de Cono de Carga.",
        "Diagrama de extracción de frames, segmentación, adaptación de máscaras, "
        "analítica geométrica y resultados.",
        width=6.15,
    )


def add_data_method(document: Document) -> None:
    document.add_heading("2. Datos de entrada y preparación", level=1)
    document.add_heading("2.1 Fuentes de entrada", level=2)
    add_body_paragraph(
        document,
        "La entrada operacional corresponde a videos o imágenes posteriores de "
        "camiones CAEX. Los videos se recorren con OpenCV y se extraen frames "
        "según un intervalo temporal configurable. Cada frame conserva su "
        "relación con el video de origen mediante manifests."
    )
    add_body_paragraph(
        document,
        "Durante el desarrollo se utilizaron además exportaciones manuales de "
        "CVAT en formato Segmentation Mask. Las imágenes se almacenan en "
        "JPEGImages y las máscaras en SegmentationClass. La clase Carga se "
        "codifica con RGB [102, 255, 102] y el fondo con [0, 0, 0]."
    )

    document.add_heading("2.2 Contrato común de máscara", level=2)
    add_body_paragraph(
        document,
        "Las máscaras manuales y las predicciones del modelo se convierten al "
        "mismo formato RGB. Esta decisión desacopla el origen de la máscara del "
        "analítico: la geometría recibe el mismo tipo de archivo tanto para "
        "validación con CVAT como para inferencia automática."
    )
    add_figure(
        document,
        "Figura_C3_cvat_mascaras.jpg",
        "Figura 2. Relación entre imagen original, anotación manual y máscara RGB.",
        "Ejemplo de imagen de un camión, anotación manual de la carga y máscara "
        "binaria exportada desde CVAT.",
        width=6.1,
    )

    document.add_heading("2.3 Supuestos de observación", level=2)
    add_bullet(
        document,
        "La cámara observa el camión desde atrás y la carga visible debe ocupar "
        "una región suficiente del frame.",
    )
    add_bullet(
        document,
        "No existe una máscara de tolva ni una ROI operacional en la versión "
        "actual.",
    )
    add_bullet(
        document,
        "Las mediciones son relativas a la propia carga segmentada y no a la "
        "posición física del camión.",
    )
    add_bullet(
        document,
        "Frames consecutivos pueden representar el mismo evento y no deben "
        "interpretarse automáticamente como observaciones independientes.",
    )


def add_perception_method(document: Document) -> None:
    document.add_heading("3. Segmentación automática de la carga", level=1)
    add_body_paragraph(
        document,
        "La percepción operacional utiliza un modelo YOLOv8m-seg entrenado para "
        "segmentación de instancias. El modelo contiene las clases caex y "
        "material; el pipeline selecciona material mediante class_id=1."
    )

    document.add_heading("3.1 Inferencia y selección de instancia", level=2)
    add_body_paragraph(
        document,
        "La inferencia se ejecuta localmente con PyTorch y Ultralytics. Para "
        "cada frame se filtran detecciones con confianza mínima 0,35 y se "
        "conserva la instancia de material con mayor área. El modo largest "
        "reduce el riesgo de incorporar acopios exteriores cuando no existe una "
        "ROI de tolva."
    )
    add_body_paragraph(
        document,
        "La configuración reproducible utilizada por Cono emplea imgsz=640, "
        "class_id=1 y modo largest. El modelo es común con Granulometría, pero "
        "los parámetros de inferencia y el postprocesamiento pertenecen a cada "
        "consumidor."
    )

    document.add_heading("3.2 Adaptación a máscara analítica", level=2)
    add_body_paragraph(
        document,
        "Los polígonos o máscaras de instancia se rasterizan al tamaño original "
        "del frame. Los píxeles seleccionados se escriben con RGB [102, 255, "
        "102] y el resto con negro. El adaptador genera una máscara PNG por "
        "frame, un manifest CSV y un resumen JSON."
    )

    stages = document.add_table(rows=1, cols=3)
    stages.style = "Table Grid"
    stages.rows[0].cells[0].text = "Etapa"
    stages.rows[0].cells[1].text = "Entrada"
    stages.rows[0].cells[2].text = "Salida"
    rows = (
        ("Inferencia", "Frame RGB", "Instancias y confianza YOLO"),
        ("Filtrado", "Predicciones", "Clase material sobre el umbral"),
        ("Selección", "Instancias válidas", "Máscara de mayor área"),
        ("Rasterización", "Máscara seleccionada", "PNG RGB compatible con CVAT"),
        ("Trazabilidad", "Estado del procesamiento", "manifest.csv y summary.json"),
    )
    for values in rows:
        cells = stages.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(stages, [1800, 3500, 4060])

    add_callout(
        document,
        "Separación de responsabilidades",
        "El modelo determina qué píxeles corresponden a carga visible. El "
        "analítico no completa regiones omitidas ni corrige material segmentado "
        "fuera de la tolva; por ello la calidad de la máscara condiciona todas "
        "las mediciones posteriores.",
        fill=LIGHT_BLUE,
    )


def add_mask_processing(document: Document) -> None:
    document.add_heading("4. Conversión y control de calidad de la máscara", level=1)
    document.add_heading("4.1 Conversión binaria", level=2)
    add_body_paragraph(
        document,
        "La máscara RGB se convierte a binaria detectando el color objetivo. "
        "Los colores inesperados se contabilizan para diagnóstico. Sobre la "
        "máscara binaria se aplica apertura morfológica de 3 píxeles, cierre de "
        "5 píxeles, relleno de huecos y conservación del componente conectado "
        "principal."
    )

    document.add_heading("4.2 Validaciones geométricas", level=2)
    add_body_paragraph(
        document,
        "Antes de calcular indicadores se verifica que la región tenga al menos "
        "100 píxeles, un bounding box mínimo de 10 por 10 píxeles y un perfil "
        "superior con al menos 20 columnas válidas."
    )
    for text in (
        "Máscara ausente, vacía o demasiado pequeña.",
        "Bounding box inválido o insuficiente.",
        "Perfil superior con pocos puntos.",
        "Máscara que toca el borde izquierdo, derecho, superior o inferior del "
        "frame.",
    ):
        add_bullet(document, text)
    add_body_paragraph(
        document,
        "Los casos que incumplen estas condiciones se conservan en el CSV con "
        "predicted_class=no_evaluable, is_valid=False y un error_reason "
        "explícito. El descarte de bordes evita interpretar como pico una parte "
        "de la carga que quedó fuera de la imagen."
    )

    document.add_heading("4.3 Referencia interna", level=2)
    add_body_paragraph(
        document,
        "El bounding box es el rectángulo mínimo que contiene la máscara limpia. "
        "Su centro divide la región en mitades izquierda y derecha y actúa como "
        "referencia para normalizar posiciones. El bounding box no representa "
        "la tolva: si la segmentación es incompleta, la referencia también se "
        "desplaza."
    )
    add_figure(
        document,
        "Figura_C5_elementos_geometricos.jpg",
        "Figura 3. Elementos geométricos calculados sobre la carga visible.",
        "Overlay de una carga segmentada con bounding box, eje interno, "
        "centroide, perfil superior y pico relativo.",
        width=6.1,
    )


def add_geometric_method(document: Document) -> None:
    document.add_heading("5. Analítica geométrica por frame", level=1)
    add_body_paragraph(
        document,
        "La descripción geométrica se construye exclusivamente sobre la máscara "
        "validada. Los indicadores principales son la posición relativa del "
        "pico, la simetría interna de área y la posición relativa del centroide."
    )

    document.add_heading("5.1 Centroide y normalización", level=2)
    add_body_paragraph(
        document,
        "El centroide se obtiene mediante momentos binarios. Su coordenada "
        "horizontal se normaliza respecto del centro y ancho del bounding box:"
    )
    add_fraction_equation(
        document,
        [("c", "x,norm")],
        [("c", "x"), " - ", ("b", "x")],
        [("w", "bbox")],
    )
    add_body_paragraph(
        document,
        "El centroide es un indicador descriptivo de confirmación. No aporta un "
        "voto independiente a la clasificación vigente."
    )

    document.add_heading("5.2 Perfil superior y pico relativo", level=2)
    add_body_paragraph(
        document,
        "Para cada columna con carga se selecciona el primer píxel desde arriba. "
        "La secuencia y_top(x) representa el contorno visible superior. El perfil "
        "se suaviza con Savitzky-Golay usando ventana 31 y polinomio de orden 2; "
        "si SciPy no está disponible, se utiliza una media móvil de 15 puntos."
    )
    add_body_paragraph(
        document,
        "El pico principal corresponde al mínimo vertical del perfil suavizado. "
        "Las posiciones que pertenecen a una meseta próxima al ápice, dentro de "
        "una tolerancia de 6 píxeles, se colapsan hacia su centro para reducir "
        "picos falsos."
    )
    add_fraction_equation(
        document,
        [("p", "x,norm")],
        [("p", "x"), " - ", ("b", "x")],
        [("w", "bbox")],
    )

    document.add_heading("5.3 Simetría interna de área", level=2)
    add_body_paragraph(
        document,
        "La máscara se divide por el centro del bounding box. La diferencia "
        "normalizada entre ambas mitades se calcula como:"
    )
    add_fraction_equation(
        document,
        [("S", "A")],
        [("A", "R"), " - ", ("A", "L")],
        [("A", "R"), " + ", ("A", "L")],
    )
    add_body_paragraph(
        document,
        "Valores negativos indican predominio visual interno hacia la izquierda; "
        "valores positivos, hacia la derecha; valores próximos a cero, una "
        "distribución aproximadamente simétrica."
    )

    document.add_heading("5.4 Indicadores diagnósticos", level=2)
    add_fraction_equation(
        document,
        [("R", "p")],
        ["std(", ("y", "raw"), " - ", ("y", "smooth"), ")"],
        [("h", "bbox")],
    )
    add_body_paragraph(
        document,
        "La irregularidad del perfil y el número de picos relevantes se guardan "
        "como diagnóstico. No generan las clases irregular ni doble_pico en la "
        "versión actual."
    )


def add_rules_method(document: Document) -> None:
    document.add_heading("6. Clasificación visual por reglas", level=1)
    add_body_paragraph(
        document,
        "La decisión combina dos señales laterales normalizadas. Cada señal "
        "aporta un voto de peso 0,50 al lado correspondiente cuando supera su "
        "umbral. Si ninguna señal supera los umbrales, el frame se clasifica "
        "como centrado."
    )

    rules = document.add_table(rows=1, cols=4)
    rules.style = "Table Grid"
    for cell, value in zip(
        rules.rows[0].cells,
        ("Indicador", "Condición izquierda", "Zona central", "Condición derecha"),
    ):
        cell.text = value
    rule_rows = (
        (
            "Pico relativo",
            "peak_x_norm < -0,06",
            "-0,06 a 0,06",
            "peak_x_norm > 0,06",
        ),
        (
            "Simetría de área",
            "area_symmetry < -0,07",
            "-0,07 a 0,07",
            "area_symmetry > 0,07",
        ),
    )
    for values in rule_rows:
        cells = rules.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(rules, [2200, 2500, 2160, 2500])

    document.add_heading("6.1 Resolución de la clase", level=2)
    for text in (
        "Sin votos laterales: centrado.",
        "Puntaje izquierdo mayor: izquierda.",
        "Puntaje derecho mayor: derecha.",
        "Máscara inválida o evidencia insuficiente: no_evaluable.",
    ):
        add_bullet(document, text)
    add_body_paragraph(
        document,
        "Con pesos iguales, señales contradictorias pueden producir empate. La "
        "implementación vigente lo resuelve hacia la izquierda debido a la "
        "comparación left_score >= right_score. Este comportamiento se mantiene "
        "documentado como limitación y debe revisarse en una calibración futura."
    )

    document.add_heading("6.2 Interpretación", level=2)
    add_callout(
        document,
        "Interpretación permitida",
        "La clase describe una tendencia visual interna de la forma segmentada. "
        "No corresponde al centro de masa real, al balance físico del camión, al "
        "desplazamiento respecto de la tolva ni a la distribución real del "
        "payload.",
        fill=LIGHT_BLUE,
    )

    document.add_heading("6.3 Variables registradas", level=2)
    features = document.add_table(rows=1, cols=2)
    features.style = "Table Grid"
    features.rows[0].cells[0].text = "Grupo"
    features.rows[0].cells[1].text = "Variables principales"
    rows = (
        ("Máscara", "mask_area_px, unexpected_pixel_count, unexpected_color_count"),
        (
            "Bounding box",
            "bbox_xmin, bbox_ymin, bbox_xmax, bbox_ymax, bbox_width, bbox_height, "
            "bbox_aspect_ratio",
        ),
        (
            "Control de borde",
            "touches_left_border, touches_right_border, touches_top_border, "
            "touches_bottom_border, is_truncated_by_frame",
        ),
        ("Centroide", "centroid_x, centroid_y, centroid_x_norm"),
        ("Perfil y pico", "peak_x, peak_y, peak_x_norm, profile_roughness, num_peaks"),
        ("Área", "area_left, area_right, area_symmetry"),
        ("Decisión", "predicted_class, is_valid, error_reason, warnings"),
    )
    for values in rows:
        cells = features.add_row().cells
        cells[0].text = values[0]
        cells[1].text = values[1]
    format_table(features, [2200, 7160])


def add_pipeline_method(document: Document) -> None:
    document.add_heading("7. Orquestación y agregación temporal", level=1)
    add_body_paragraph(
        document,
        "El orquestador recibe uno o más videos locales y ejecuta todas las "
        "etapas con un nombre de corrida. La relación entre video, frame, máscara "
        "y resultado se conserva mediante manifests."
    )
    stages = (
        "Descubrir videos y registrar sus metadatos.",
        "Extraer frames con el intervalo configurado.",
        "Ejecutar segmentación y producir máscaras compatibles.",
        "Aplicar el analítico geométrico a cada par imagen-máscara.",
        "Generar overlays, features y estadísticas.",
        "Agrupar frames por video y construir un veredicto temporal.",
    )
    for text in stages:
        paragraph = document.add_paragraph(style="List Number")
        run = paragraph.add_run(text)
        set_run_font(run, color=INK)

    add_figure(
        document,
        "Figura_C6_secuencia_pipeline.jpg",
        "Figura 4. Secuencia de ejecución del pipeline completo.",
        "Diagrama de secuencia desde video y extracción de frames hasta "
        "segmentación, analítica y resultados.",
        width=6.1,
    )

    document.add_heading("7.1 Veredicto por video", level=2)
    add_body_paragraph(
        document,
        "Primero se exige que al menos el 50 % de los frames extraídos sea "
        "evaluable. Luego se calculan las proporciones de centrado, izquierda y "
        "derecha. Si una clase alcanza el 60 %, se emite el veredicto "
        "correspondiente. En ausencia de mayoría se revisan las medianas del "
        "pico relativo y la simetría de área; si las señales son insuficientes "
        "o contradictorias, el resultado es mixto_o_inestable."
    )
    add_body_paragraph(
        document,
        "La confianza informada corresponde a la estabilidad interna de las "
        "clases observadas y no a una probabilidad calibrada del sistema."
    )


def add_outputs_validation(document: Document) -> None:
    document.add_heading("8. Salidas, revisión y validación", level=1)
    document.add_heading("8.1 Productos por corrida", level=2)
    outputs = document.add_table(rows=1, cols=3)
    outputs.style = "Table Grid"
    for cell, value in zip(outputs.rows[0].cells, ("Salida", "Contenido", "Uso")):
        cell.text = value
    rows = (
        ("features.csv", "Variables por frame", "Análisis cuantitativo y auditoría"),
        ("overlays/", "Máscara, bbox, perfil y clase", "Revisión visual"),
        ("video_summary.csv", "Conteos, ratios y veredicto", "Resumen por video"),
        ("pipeline_summary.json", "Rutas, tiempos y estados", "Trazabilidad"),
        ("indicator_summary.csv", "Medianas y cuartiles", "Descripción de indicadores"),
        ("indicator_distributions.csv", "Histogramas de 20 intervalos", "Análisis estadístico"),
    )
    for values in rows:
        cells = outputs.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(outputs, [2600, 3500, 3260])

    document.add_heading("8.2 Estrategia de validación", level=2)
    add_body_paragraph(
        document,
        "La validación del método se organizó en tres niveles complementarios:"
    )
    add_bullet(
        document,
        "Máscaras manuales de CVAT para comprobar el analítico sin depender del "
        "modelo de percepción.",
    )
    add_bullet(
        document,
        "Prueba técnica balanceada de 90 imágenes, con 30 ejemplos visuales por "
        "clase, para construir matriz de confusión, exactitud, precisión, "
        "sensibilidad y F1.",
    )
    add_bullet(
        document,
        "Muestra descriptiva de 1.248 imágenes para observar evaluabilidad, "
        "distribución de clases y causas de descarte en un volumen mayor.",
    )
    add_body_paragraph(
        document,
        "La prueba balanceada mide el resultado extremo a extremo del pipeline "
        "en una muestra seleccionada. La muestra de 1.248 imágenes no posee "
        "ground truth independiente para cada frame y, por tanto, describe "
        "salidas pero no exactitud."
    )

    document.add_heading("8.3 Revisión visual", level=2)
    add_body_paragraph(
        document,
        "Los overlays permiten verificar si la máscara cubre la carga, si toca "
        "bordes, si el perfil superior sigue el contorno y si el pico y la "
        "simetría son coherentes. Esta revisión es necesaria porque un bounding "
        "box correcto sobre una máscara incompleta sigue produciendo indicadores "
        "geométricos sesgados."
    )


def add_reproducibility(document: Document) -> None:
    document.add_heading("9. Reproducibilidad, limitaciones y uso previsto", level=1)
    document.add_heading("9.1 Reproducibilidad", level=2)
    add_body_paragraph(
        document,
        "Los umbrales geométricos se almacenan en "
        "configs/analytic_config.yaml y los parámetros de percepción en "
        "perception/configs/perception_config.yaml. La inferencia operacional se "
        "ejecuta en el entorno Python local. Docker se reserva para preparación "
        "de datasets, entrenamiento, validación y registro de experimentos en "
        "MLflow."
    )
    add_body_paragraph(
        document,
        "Cada corrida registra rutas, nombre del modelo, clase seleccionada, "
        "confianza, resolución, modo de instancia, errores, tiempos y archivos "
        "generados. Esto permite reproducir y comparar ejecuciones sin depender "
        "de la interfaz web."
    )

    document.add_heading("9.2 Limitaciones metodológicas", level=2)
    for text in (
        "No existe una referencia explícita de la tolva para medir desplazamiento "
        "respecto del camión.",
        "La imagen es bidimensional y no estima volumen, masa, densidad ni "
        "estabilidad física.",
        "La clasificación depende de la cobertura y calidad de la máscara.",
        "Las reglas fueron calibradas sobre muestras locales y requieren "
        "validación adicional en otras cámaras, fechas y condiciones.",
        "Los frames de una misma secuencia tienen dependencia temporal.",
        "Los empates entre señales laterales necesitan una regla operacional más "
        "robusta.",
    ):
        add_bullet(document, text)

    document.add_heading("9.3 Uso previsto", level=2)
    add_callout(
        document,
        "Uso recomendado",
        "La salida debe emplearse como apoyo visual y analítico para priorizar "
        "revisión, comparar tendencias y estudiar la forma visible de la carga. "
        "No reemplaza inspecciones de seguridad, mediciones físicas ni criterios "
        "operacionales definidos por especialistas.",
    )

    document.add_heading("9.4 Resumen del método", level=2)
    add_body_paragraph(
        document,
        "En síntesis, el sistema transforma videos en frames, segmenta la carga "
        "visible, normaliza la máscara a un contrato común, descarta evidencia "
        "insuficiente, calcula geometría interna, clasifica cada frame mediante "
        "reglas trazables y agrega los resultados para emitir un veredicto por "
        "video. La modularidad permite mejorar la percepción, las reglas o la "
        "agregación temporal sin cambiar el contrato entre componentes."
    )

    add_body_paragraph(
        document,
        "Fuentes técnicas internas: configuraciones en "
        "configs/analytic_config.yaml y "
        "perception/configs/perception_config.yaml. Orquestación y adaptación: "
        "scripts/run_pipeline_from_videos.py y "
        "perception/scripts/predict_to_analytic_masks.py. Implementación del "
        "analítico: src/cone_load/. Contexto consolidado: "
        "docs/contexto_completo_proyecto_cono.md."
    )


def build_document() -> Document:
    document = Document()
    configure_styles(document)
    configure_sections(document)
    add_cover(document)
    add_scope(document)
    add_data_method(document)
    add_perception_method(document)
    add_mask_processing(document)
    add_geometric_method(document)
    add_rules_method(document)
    add_pipeline_method(document)
    add_outputs_validation(document)
    add_reproducibility(document)
    return document


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(f"Documento creado: {OUTPUT}")


if __name__ == "__main__":
    main()
