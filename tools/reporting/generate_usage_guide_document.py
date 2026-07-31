from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from generate_methodology_document import (
    ACCENT,
    BLUE,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    TABLE_WIDTH_DXA,
    add_body_paragraph,
    add_bullet,
    add_callout,
    add_figure,
    add_page_number,
    format_table,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = (
    ROOT
    / "deliverables"
    / "modo_de_uso"
    / "Modo_de_uso_proyecto_cono_de_carga.docx"
)

CODE_FILL = "F4F6F9"
CODE_BORDER = "2E74B5"
GREEN_FILL = "EAF4EE"
GOLD_FILL = "FFF4D6"
RED_FILL = "FDECEC"


def configure_styles(document: Document) -> None:
    """Apply the compact_reference_guide preset."""
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Title", 28, BLUE, 0, 8),
        ("Subtitle", 13, MUTED, 0, 16),
        ("Heading 1", 16, ACCENT, 18, 10),
        ("Heading 2", 13, ACCENT, 14, 7),
        ("Heading 3", 12, BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = INK
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.color.rgb = MUTED
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)


def configure_sections(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("GUÍA OPERATIVA  |  CONO DE CARGA")
        set_run_font(run, size=8.5, color=MUTED, bold=True)

        footer = section.footer
        footer.is_linked_to_previous = False
        add_page_number(footer.paragraphs[0])


def add_code_block(document: Document, code: str, label: str | None = None) -> None:
    if label:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(label)
        set_run_font(run, size=9.5, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.10)
    paragraph.paragraph_format.right_indent = Inches(0.02)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True

    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CODE_FILL)
    properties.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), CODE_BORDER)
    borders.append(left)
    properties.append(borders)

    run = paragraph.add_run(code.strip())
    set_run_font(run, name="Consolas", size=8.4, color=INK)


def new_numbering_instance(document: Document) -> int:
    """Create a fresh Word numbering instance based on List Number."""
    numbering = document.part.numbering_part.element
    style_num_id = int(
        document.styles["List Number"]
        ._element.pPr.numPr.numId.get(qn("w:val"))
    )
    base_num = numbering.find(f'./w:num[@w:numId="{style_num_id}"]', numbering.nsmap)
    if base_num is None:
        raise RuntimeError("No se encontró la numeración base de List Number.")
    abstract = base_num.find(qn("w:abstractNumId"))
    if abstract is None:
        raise RuntimeError("List Number no referencia una numeración abstracta.")

    current_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    new_id = max(current_ids, default=0) + 1
    node = OxmlElement("w:num")
    node.set(qn("w:numId"), str(new_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), abstract.get(qn("w:val")))
    node.append(reference)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    node.append(override)
    numbering.append(node)
    return new_id


def add_numbered_item(
    document: Document,
    title: str,
    detail: str,
    *,
    num_id: int,
) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.keep_with_next = True
    properties = paragraph._p.get_or_add_pPr()
    number_properties = properties.find(qn("w:numPr"))
    if number_properties is None:
        number_properties = OxmlElement("w:numPr")
        properties.append(number_properties)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), str(num_id))
    number_properties.append(level)
    number_properties.append(number_id)
    first = paragraph.add_run(title)
    set_run_font(first, color=INK, bold=True)
    rest = paragraph.add_run(f" {detail}")
    set_run_font(rest, color=INK)


def add_key_value_table(
    document: Document,
    headers: tuple[str, ...],
    rows: tuple[tuple[str, ...], ...],
    widths: list[int],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    format_table(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(document: Document) -> None:
    document.add_paragraph().paragraph_format.space_after = Pt(28)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(26)
    run = kicker.add_run("PROYECTO MEMORIA  |  GUÍA OPERATIVA")
    set_run_font(run, size=10.5, color=ACCENT, bold=True)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Modo de uso del proyecto Cono de Carga")
    set_run_font(run, size=28, color=BLUE, bold=True)

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Continuidad de datos, entrenamiento de segmentación, validación y "
        "ejecución del pipeline"
    )
    set_run_font(run, size=13, color=MUTED)

    document.add_paragraph().paragraph_format.space_after = Pt(16)
    add_callout(
        document,
        "Propósito",
        "Servir como manual reproducible para incorporar nuevas exportaciones "
        "de CVAT, entrenar candidatos YOLO, compararlos con el modelo vigente y "
        "ejecutar la solución completa sin confundir los entornos de desarrollo "
        "y operación.",
        fill=GREEN_FILL,
    )

    document.add_paragraph().paragraph_format.space_after = Pt(12)
    metadata = document.add_table(rows=4, cols=2)
    metadata.style = "Table Grid"
    values = (
        ("Repositorio", r"D:\Memoria\Cono"),
        ("Modelo operacional", r"perception\models\primary\psinet_carga_visible.pt"),
        ("Analítico", "Clasificación visual interna: izquierda, centrado, derecha o no_evaluable"),
        ("Versión", "Julio de 2026"),
    )
    for row, (label, value) in zip(metadata.rows, values):
        row.cells[0].text = label
        row.cells[1].text = value
    format_table(metadata, [2500, 6860])

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(30)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(
        "Docker se usa para preparar datos, entrenar y validar. La inferencia, "
        "el pipeline y la web app se ejecutan localmente."
    )
    set_run_font(run, size=10, color=MUTED, italic=True)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_orientation(document: Document) -> None:
    document.add_heading("1. Orientación rápida", level=1)
    add_body_paragraph(
        document,
        "El proyecto recibe videos o imágenes de camiones CAEX vistos desde "
        "atrás. Un modelo YOLO segmenta la carga visible; un adaptador convierte "
        "la predicción en una máscara RGB; el analítico calcula indicadores "
        "relativos a la propia carga y genera resultados por frame y por video."
    )
    add_figure(
        document,
        "Figura_C2_arquitectura.jpg",
        "Figura 1. Componentes principales y contratos del sistema.",
        "Arquitectura del proyecto: entrada, segmentación, máscara, analítico y resultados.",
        width=6.1,
    )

    document.add_heading("1.1 Dos entornos con responsabilidades distintas", level=2)
    add_key_value_table(
        document,
        ("Entorno", "Se utiliza para", "No se necesita para"),
        (
            (
                "Docker / perception",
                "Preparar el dataset YOLO, entrenar, validar y registrar experimentos en MLflow.",
                "Ejecutar el pipeline, la web app o inferencia operacional.",
            ),
            (
                "Python local / web_app\\.venv",
                "Inferencia, pipeline desde videos, analítico, evaluaciones y Streamlit.",
                "Entrenar nuevos modelos dentro del flujo reproducible del repositorio.",
            ),
        ),
        [2100, 3630, 3630],
    )

    document.add_heading("1.2 Rutas que gobiernan el sistema", level=2)
    for text in (
        r"Configuración de percepción: perception\configs\perception_config.yaml.",
        r"Configuración del analítico: configs\analytic_config.yaml.",
        r"Modelo vigente: perception\models\primary\psinet_carga_visible.pt.",
        r"Modelo alternativo histórico: perception\models\alternative\best.pt.",
        r"Resultados de pipeline: outputs\pipeline_runs\<run_name>\.",
        r"Resultados de entrenamiento: perception\runs\train\<run_name>\ y perception\models\<run_name>\.",
    ):
        add_bullet(document, text)

    add_callout(
        document,
        "Alcance de los resultados",
        "Las clases describen la forma visible y la asimetría visual interna de "
        "la carga segmentada. No representan centro de masa real, balance físico "
        "del camión ni distribución real del payload.",
        fill=GOLD_FILL,
    )


def add_prerequisites(document: Document) -> None:
    document.add_heading("2. Preparación inicial", level=1)
    document.add_heading("2.1 Requisitos del equipo", level=2)
    for text in (
        "Windows con PowerShell.",
        "Python 3.12 para el entorno local del pipeline.",
        "Docker Desktop con backend WSL2 para entrenamiento.",
        "Driver NVIDIA y GPU compatible; la GPU debe aparecer en nvidia-smi.",
        "Espacio libre suficiente para imágenes, datasets YOLO, corridas y pesos.",
    ):
        add_bullet(document, text)

    add_code_block(document, "nvidia-smi", "Comprobar la GPU en el host")

    document.add_heading("2.2 Crear o reparar el entorno local", level=2)
    add_body_paragraph(
        document,
        "El lanzador de la web app puede crear este entorno automáticamente. "
        "Para controlarlo manualmente, ejecute:"
    )
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono
py -3.12 -m venv web_app\.venv
.\web_app\.venv\Scripts\Activate.ps1
python -m pip install --upgrade pip
pip install -r requirements\runtime.txt""",
        "Entorno local",
    )
    add_code_block(
        document,
        r"""python -c "import torch; print(torch.cuda.is_available()); print(torch.cuda.get_device_name(0) if torch.cuda.is_available() else 'CPU')""",
        "Comprobar PyTorch y CUDA local",
    )

    document.add_heading("2.3 Construir el contenedor de entrenamiento", level=2)
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono\perception
docker compose build
docker compose up -d
docker compose exec cone-yolo python scripts/check_env.py""",
        "Primera ejecución",
    )
    add_body_paragraph(
        document,
        "No es necesario reconstruir por cada experimento: configuración, "
        "scripts, datos, corridas y modelos están montados como volúmenes. "
        "Reconstruya únicamente al cambiar Dockerfile, requirements/docker.txt "
        "o la base CUDA/PyTorch."
    )


def add_data_workflow(document: Document) -> None:
    document.add_heading("3. Incorporar más datos", level=1)
    add_callout(
        document,
        "Regla principal",
        "La separación train/val/test debe realizarse por video, evento, fecha o "
        "secuencia antes de preparar el dataset. Frames consecutivos de un mismo "
        "paso de camión no deben quedar repartidos entre conjuntos, porque eso "
        "infla artificialmente las métricas.",
        fill=GOLD_FILL,
    )

    document.add_heading("3.1 Planificar la ampliación", level=2)
    plan_numbering = new_numbering_instance(document)
    for title, detail in (
        (
            "Definir el objetivo.",
            "Especifique qué error se busca reducir: iluminación, polvo, distancia, cámara, oclusión, carga parcial o una geometría poco representada.",
        ),
        (
            "Evitar duplicados cercanos.",
            "Muestree con separación temporal suficiente; cientos de frames casi idénticos aportan menos que escenarios nuevos.",
        ),
        (
            "Conservar un test congelado.",
            "No use el conjunto test para seleccionar épocas, umbrales ni hiperparámetros.",
        ),
        (
            "Registrar procedencia.",
            "Anote videos, fechas, cámaras, criterio de muestreo, responsable del etiquetado y versión del labelmap.",
        ),
    ):
        add_numbered_item(document, title, detail, num_id=plan_numbering)

    document.add_heading("3.2 Elegir el contrato de clases antes de etiquetar", level=2)
    add_key_value_table(
        document,
        ("Camino", "Contrato", "Uso recomendado"),
        (
            (
                "Modelo compartido",
                "caex=0 y material=1",
                "Continuar el modelo operacional común con Granulometría. Requiere exportaciones y conversión compatibles con ambas clases.",
            ),
            (
                "Candidato solo Cono",
                "Carga=0",
                "Usar el preparador CVAT actual y entrenar un candidato independiente. Debe evaluarse con --class-id 0.",
            ),
        ),
        [1900, 2300, 5160],
    )
    add_callout(
        document,
        "No mezclar contratos",
        "El modelo operacional usa material=1, pero el dataset generado por la "
        "configuración actual usa Carga=0. No haga fine-tuning del modelo "
        "operacional con ese dataset de una clase sin adaptar explícitamente las "
        "clases y conservar los datos necesarios; podría perder o reinterpretar "
        "la clase caex.",
        fill=RED_FILL,
    )

    document.add_heading("3.3 Etiquetar y exportar desde CVAT", level=2)
    add_body_paragraph(
        document,
        "Para el camino de una clase incluido actualmente, etiquete únicamente "
        "la carga visible. Exporte en formato Segmentation Mask y compruebe que "
        "cada paquete tenga la estructura siguiente:"
    )
    add_code_block(
        document,
        """labelmap.txt
JPEGImages/
SegmentationClass/""",
        "Estructura esperada",
    )
    add_code_block(
        document,
        """Carga:102,255,102::
background:0,0,0::""",
        "Colores esperados en labelmap.txt",
    )
    for text in (
        "Use máscaras PNG; JPEG altera los colores y puede impedir detectar la clase.",
        "Mantenga una política consistente frente a oclusiones y truncamiento.",
        "Incluya escenas difíciles y fondos variados, no solo ejemplos limpios.",
        "Revise visualmente una muestra de cada exportación antes de entrenar.",
    ):
        add_bullet(document, text)

    document.add_heading("3.4 Combinar exportaciones divididas en ZIP", level=2)
    add_body_paragraph(
        document,
        "Deposite los ZIP en input/train, input/val e input/test del combinador. "
        "Después ejecute:"
    )
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono\perception
python .\tools\cvat_zip_merger\merge_cvat_zips.py --overwrite""",
        "Combinar ZIP",
    )
    add_body_paragraph(
        document,
        "Los conjuntos combinados quedan en tools/cvat_zip_merger/output. Copie "
        "cada split a data/raw:"
    )
    add_code_block(
        document,
        r"""Copy-Item .\tools\cvat_zip_merger\output\train .\data\raw\train -Recurse -Force
Copy-Item .\tools\cvat_zip_merger\output\val .\data\raw\val -Recurse -Force
Copy-Item .\tools\cvat_zip_merger\output\test .\data\raw\test -Recurse -Force""",
        "Copiar splits",
    )

    document.add_heading("3.5 Verificación previa", level=2)
    for text in (
        "Cada imagen tiene una máscara con el mismo nombre base.",
        "labelmap.txt coincide con el contrato de clases del experimento.",
        "Los splits no comparten secuencias ni frames vecinos.",
        "No hay máscaras vacías involuntarias ni colores inesperados.",
        "El número de muestras por split y por origen quedó registrado.",
    ):
        add_bullet(document, text)


def add_training_workflow(document: Document) -> None:
    document.add_heading("4. Preparar, entrenar y validar un candidato", level=1)
    document.add_heading("4.1 Preparar el dataset YOLO", level=2)
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono\perception
docker compose up -d
docker compose exec cone-yolo python scripts/prepare_yolo_dataset.py --overwrite""",
        "Conversión CVAT a YOLO-seg",
    )
    add_body_paragraph(
        document,
        "El script respeta data/raw/train, val y test. Antes de entrenar, revise "
        "data/yolo/carga_visible_seg/manifest.csv y dataset.yaml; el conteo debe "
        "coincidir con los datos aprobados."
    )
    add_code_block(
        document,
        r"""docker compose exec cone-yolo python scripts/prepare_yolo_dataset.py `
  --train-cvat-dir "data/raw/train" `
  --val-cvat-dir "data/raw/val" `
  --test-cvat-dir "data/raw/test" `
  --config "configs/perception_config.yaml" `
  --overwrite""",
        "Forma explícita",
    )

    document.add_heading("4.2 Configurar trazabilidad en MLflow", level=2)
    add_body_paragraph(
        document,
        "El contenedor carga perception/.env.mlflow. Para un servidor remoto, "
        "complete MLFLOW_TRACKING_URI y las credenciales del almacén de "
        "artefactos. No versione secretos. Si no se define un servidor remoto, "
        "se utiliza perception/mlruns como tracking local."
    )
    add_code_block(
        document,
        """MLFLOW_TRACKING_URI=http://<servidor>:5000
MLFLOW_S3_ENDPOINT_URL=<endpoint>
AWS_ACCESS_KEY_ID=<credencial>
AWS_SECRET_ACCESS_KEY=<credencial>
AWS_DEFAULT_REGION=<region>""",
        "Variables principales de .env.mlflow",
    )
    add_figure(
        document,
        "Figura_C4_entrenamiento_mlflow.png",
        "Figura 2. Entrenamiento y trazabilidad experimental con MLflow.",
        "Curvas de entrenamiento, métricas de segmentación y datos del experimento MLflow.",
        width=6.15,
    )

    document.add_heading("4.3 Entrenar desde arquitectura", level=2)
    add_body_paragraph(
        document,
        "Use un nombre único y descriptivo. El ejemplo siguiente entrena un "
        "candidato de una clase con arquitectura YOLOv8m-seg:"
    )
    add_code_block(
        document,
        r"""docker compose exec cone-yolo python scripts/train_yolo_seg.py `
  --config "configs/perception_config.yaml" `
  --model yolov8m-seg.yaml `
  --no-pretrained `
  --epochs 100 `
  --batch 8 `
  --imgsz 640 `
  --device 0 `
  --run-name "carga_v8m_datos_2026_08"
""",
        "Entrenamiento desde cero",
    )
    add_body_paragraph(
        document,
        "Para una prueba rápida puede usar yolov8n-seg.yaml y menos épocas. "
        "Cambie un grupo pequeño de variables por experimento para que la "
        "comparación sea interpretable."
    )

    document.add_heading("4.4 Fine-tuning: usar solo con clases compatibles", level=2)
    add_code_block(
        document,
        r"""docker compose exec cone-yolo python scripts/train_yolo_seg.py `
  --config "configs/perception_config.yaml" `
  --model "models/<modelo_base>/best.pt" `
  --pretrained `
  --epochs 50 `
  --batch 8 `
  --imgsz 640 `
  --device 0 `
  --run-name "ajuste_<descripcion>"
""",
        "Plantilla de fine-tuning",
    )
    add_callout(
        document,
        "Condición de uso",
        "El número, orden y significado de las clases del dataset deben ser "
        "compatibles con los pesos base. Si no lo son, entrene un candidato "
        "separado desde arquitectura o adapte formalmente la preparación del "
        "dataset antes de continuar.",
        fill=GOLD_FILL,
    )

    document.add_heading("4.5 Validar la segmentación", level=2)
    add_code_block(
        document,
        r"""docker compose exec cone-yolo python scripts/validate_yolo_seg.py `
  --model "models/carga_v8m_datos_2026_08/best.pt" `
  --config "configs/perception_config.yaml" `
  --device 0""",
        "Validación del candidato",
    )
    for text in (
        "Compare precisión, recall y mAP de máscara con el baseline en el mismo test congelado.",
        "Revise resultados visuales, especialmente falsos positivos, máscaras parciales y bordes.",
        "No seleccione el candidato solo por una métrica promedio; compruebe escenarios críticos.",
        "Confirme que best.pt y last.pt quedaron en perception/models/<run_name>/ y en MLflow.",
    ):
        add_bullet(document, text)

    document.add_heading("4.6 Probar el impacto en el analítico", level=2)
    add_body_paragraph(
        document,
        "Una mejora de segmentación no garantiza una mejora del veredicto. "
        "Ejecute el candidato sobre un benchmark independiente y compare la "
        "clasificación izquierda/centrado/derecha, la proporción no_evaluable, "
        "los overlays y los indicadores geométricos."
    )
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono
.\web_app\.venv\Scripts\Activate.ps1

python perception\scripts\predict_to_analytic_masks.py `
  --model "perception\models\carga_v8m_datos_2026_08\best.pt" `
  --images-dir "<benchmark_de_imagenes>" `
  --masks-dir "perception\model_testing\analytic_masks\carga_v8m_datos_2026_08" `
  --class-id 0 `
  --conf 0.35 `
  --device 0 `
  --overwrite""",
        "Generar máscaras con un candidato Carga=0",
    )
    add_code_block(
        document,
        r"""python scripts\run_simple_analytic.py `
  --images-dir "<benchmark_de_imagenes>" `
  --masks-dir "perception\model_testing\analytic_masks\carga_v8m_datos_2026_08" `
  --output-dir "outputs\evaluaciones\carga_v8m_datos_2026_08" `
  --config "configs\analytic_config.yaml"
""",
        "Analítico sobre el candidato",
    )

    document.add_heading("4.7 Promover un modelo sin perder trazabilidad", level=2)
    promotion_numbering = new_numbering_instance(document)
    for title, detail in (
        ("Conservar el candidato.", "No sobrescriba el modelo vigente durante la experimentación."),
        ("Guardar evidencia.", "Archive configuración, manifest, métricas, overlays, best.pt y run de MLflow."),
        ("Comparar con baseline.", "Use el mismo test y los mismos parámetros de inferencia."),
        ("Decidir el class_id.", "Use 1 para material en el modelo compartido o 0 para un candidato Carga de una clase."),
        ("Actualizar configuración.", "Cambie prediction.model y prediction.class_id de forma explícita y documentada."),
        ("Ejecutar prueba de humo.", "Corra un video conocido, revise CSV, JSON, overlays y web app."),
        ("Versionar el cambio.", "Haga commit de configuración y documentación; no agregue datasets ni corridas masivas."),
    ):
        add_numbered_item(document, title, detail, num_id=promotion_numbering)


def add_operation(document: Document) -> None:
    document.add_heading("5. Ejecutar el proyecto", level=1)
    document.add_heading("5.1 Pipeline completo desde videos", level=2)
    add_body_paragraph(
        document,
        "Coloque uno o más videos en data/pipeline_input/videos. Cada video se "
        "procesa por separado y recibe un veredicto independiente."
    )
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono
.\web_app\.venv\Scripts\Activate.ps1

python scripts\run_pipeline_from_videos.py `
  --videos-dir "data\pipeline_input\videos" `
  --run-name "prueba_video" `
  --interval-seconds 0.5 `
  --max-frames-per-video 20 `
  --conf 0.35 `
  --device 0 `
  --overwrite""",
        "Pipeline operacional",
    )
    add_figure(
        document,
        "Figura_C6_secuencia_pipeline.jpg",
        "Figura 3. Secuencia operacional del pipeline desde video hasta resultados.",
        "Flujo secuencial de extracción, segmentación, analítico y veredicto por video.",
        width=6.15,
    )
    add_body_paragraph(
        document,
        "Use --interval-seconds 0 para extraer todos los frames. Valores menores "
        "producen un análisis temporal más denso, pero aumentan tiempo, memoria y "
        "redundancia. Los resultados quedan en outputs/pipeline_runs/<run_name>/."
    )

    document.add_heading("5.2 Aplicación web", level=2)
    add_body_paragraph(
        document,
        "En Windows, haga doble clic en Abrir web app Cono de Carga.bat. El "
        "lanzador reutiliza web_app/.venv y abre la interfaz en "
        "http://127.0.0.1:8502. No necesita Docker."
    )
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono
.\web_app\.venv\Scripts\Activate.ps1
streamlit run web_app\app.py --server.address 127.0.0.1 --server.port 8502""",
        "Inicio manual de la web app",
    )

    document.add_heading("5.3 Inferencia de segmentación sobre imágenes", level=2)
    add_code_block(
        document,
        r"""python perception\scripts\predict_yolo_seg.py `
  --model "perception\models\primary\psinet_carga_visible.pt" `
  --source "<carpeta_de_imagenes>" `
  --config "perception\configs\perception_config.yaml" `
  --conf 0.35 `
  --device 0""",
        "Predicción visual local",
    )

    document.add_heading("5.4 Analítico con máscaras manuales o existentes", level=2)
    add_code_block(
        document,
        r"""python scripts\run_simple_analytic.py `
  --images-dir "<ruta_a_JPEGImages>" `
  --masks-dir "<ruta_a_SegmentationClass>" `
  --output-dir "outputs\analitico_manual" `
  --config "configs\analytic_config.yaml""",
        "Solo analítico",
    )
    add_body_paragraph(
        document,
        "Esta modalidad es útil para validar reglas geométricas sin involucrar "
        "el error del modelo de percepción. Genera features.csv y un overlay por "
        "imagen."
    )


def add_commands(document: Document) -> None:
    document.add_heading("6. Comandos útiles", level=1)
    document.add_heading("6.1 Docker y GPU", level=2)
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono\perception
docker compose up -d
docker compose ps
docker compose exec cone-yolo python scripts/check_env.py
docker compose logs --tail 100 cone-yolo
docker compose down""",
        "Ciclo normal",
    )
    add_code_block(
        document,
        r"""docker compose down
docker compose build --no-cache
docker compose up -d""",
        "Reconstruir después de cambiar dependencias",
    )
    add_code_block(
        document,
        r"""docker compose --profile ui up -d mlflow-ui""",
        "MLflow local en http://localhost:5000",
    )

    document.add_heading("6.2 Ayuda incorporada", level=2)
    add_code_block(
        document,
        r"""python scripts\run_pipeline_from_videos.py --help
python scripts\run_simple_analytic.py --help
python perception\scripts\predict_to_analytic_masks.py --help
docker compose exec cone-yolo python scripts/prepare_yolo_dataset.py --help
docker compose exec cone-yolo python scripts/train_yolo_seg.py --help
docker compose exec cone-yolo python scripts/validate_yolo_seg.py --help""",
        "Consultar parámetros",
    )

    document.add_heading("6.3 Pruebas y diagnóstico local", level=2)
    add_code_block(
        document,
        r"""cd D:\Memoria\Cono
.\web_app\.venv\Scripts\Activate.ps1
python -m pytest tests -q
python -c "import cv2, numpy, pandas, scipy, torch, ultralytics; print('dependencias OK')"
git status --short""",
        "Prueba de humo del repositorio",
    )

    document.add_heading("6.4 Parámetros frecuentes del pipeline", level=2)
    add_key_value_table(
        document,
        ("Parámetro", "Efecto"),
        (
            ("--interval-seconds", "Separación temporal entre frames; 0 extrae todos."),
            ("--max-frames-per-video", "Límite para pruebas rápidas; omítalo para no limitar."),
            ("--recursive", "Busca videos dentro de subdirectorios."),
            ("--conf", "Umbral de confianza de segmentación; valor operacional 0.35."),
            ("--device", "0 para la primera GPU; cpu para ejecutar sin CUDA."),
            ("--class-id", "1 para material del modelo operacional; 0 para candidato Carga."),
            ("--mask-mode", "largest conserva la instancia principal; merge_all une instancias."),
            ("--overwrite", "Reemplaza una corrida previa con el mismo nombre."),
        ),
        [2600, 6760],
    )

    document.add_heading("6.5 PowerShell sin errores de continuación", level=2)
    add_body_paragraph(
        document,
        "El carácter grave (`) debe ser el último carácter de la línea. No copie "
        "los símbolos >> mostrados por PowerShell: son el prompt de continuación, "
        "no parte del comando. También puede ejecutar el comando completo en una "
        "sola línea."
    )


def add_outputs_and_troubleshooting(document: Document) -> None:
    document.add_heading("7. Salidas, respaldo y solución de problemas", level=1)
    document.add_heading("7.1 Qué conservar de cada experimento", level=2)
    for text in (
        "Nombre y objetivo del experimento.",
        "Versión o commit del código y configuración YAML usada.",
        "Manifest del dataset, procedencia y separación de splits.",
        "Métricas de validación y evidencia visual de errores.",
        "best.pt, last.pt y enlace o run_id de MLflow.",
        "Resultados end-to-end sobre el benchmark independiente.",
        "Decisión final: rechazado, candidato o promovido, con justificación.",
    ):
        add_bullet(document, text)

    document.add_heading("7.2 Directorios de salida", level=2)
    add_key_value_table(
        document,
        ("Salida", "Ubicación"),
        (
            ("Dataset YOLO", "perception\\data\\yolo\\carga_visible_seg\\"),
            ("Corrida Ultralytics", "perception\\runs\\train\\<run_name>\\"),
            ("Pesos consolidados", "perception\\models\\<run_name>\\"),
            ("Tracking MLflow local", "perception\\mlruns\\"),
            ("Pipeline por video", "outputs\\pipeline_runs\\<run_name>\\"),
            ("Analítico directo", "<output-dir>\\tables\\features.csv y <output-dir>\\overlays\\"),
            ("Web app", "web_app\\outputs\\runs\\<run_name>\\"),
        ),
        [2900, 6460],
    )

    document.add_heading("7.3 Problemas frecuentes", level=2)
    add_key_value_table(
        document,
        ("Síntoma", "Causa probable y acción"),
        (
            (
                "docker compose: no configuration file",
                r"Ejecute el comando desde D:\Memoria\Cono\perception.",
            ),
            (
                "torch.cuda.is_available() = False",
                "Revise nvidia-smi, driver, soporte GPU de Docker o instalación local de PyTorch.",
            ),
            (
                "ModuleNotFoundError: boto3",
                "La imagen se construyó antes de agregar la dependencia; reconstruya con --no-cache.",
            ),
            (
                "Máscaras vacías o todos no_evaluable",
                "Verifique class_id, confianza, formato PNG, labelmap, color RGB y predicción visual.",
            ),
            (
                "Carga descartada por borde",
                "El analítico rechaza máscaras que tocan cualquier borde porque la forma visible está truncada.",
            ),
            (
                "Picos falsos o máscara parcial",
                "Inspeccione el overlay; el problema puede provenir de segmentación incompleta o ruido del contorno.",
            ),
            (
                "No aparece best.pt",
                r"Revise perception\runs\train\<run_name>\weights y perception\models\<run_name>.",
            ),
            (
                "El comando muestra errores con >>",
                "Pegó el prompt de continuación; vuelva a ejecutar sin esos símbolos.",
            ),
        ),
        [3050, 6310],
    )

    document.add_heading("7.4 Cierre de sesión", level=2)
    add_code_block(
        document,
        r"""# Salir del entorno local
deactivate

# Detener el contenedor de entrenamiento
cd D:\Memoria\Cono\perception
docker compose down""",
        "Detener entornos",
    )


def add_final_checklists(document: Document) -> None:
    document.add_heading("8. Listas de control", level=1)
    document.add_heading("8.1 Antes de entrenar", level=2)
    for text in (
        "[ ] El objetivo y el baseline están definidos.",
        "[ ] El contrato de clases coincide con el modelo o arquitectura elegida.",
        "[ ] Train, val y test están separados por origen temporal.",
        "[ ] El test permanece congelado.",
        "[ ] Las máscaras y labelmap fueron revisados visualmente.",
        "[ ] manifest.csv y dataset.yaml tienen conteos y clases correctos.",
        "[ ] La GPU es visible y MLflow registra el experimento.",
        "[ ] El run-name es único y descriptivo.",
    ):
        add_bullet(document, text)

    document.add_heading("8.2 Antes de promover", level=2)
    for text in (
        "[ ] El candidato mejora o mantiene métricas en el test congelado.",
        "[ ] Los errores críticos fueron inspeccionados en overlays.",
        "[ ] La evaluación end-to-end no empeora clasificación ni no_evaluable.",
        "[ ] El class_id operacional está documentado.",
        "[ ] Los pesos y artefactos están respaldados en MLflow y models/<run_name>.",
        "[ ] Se ejecutó un video de prueba completo y la web app abre sus resultados.",
        "[ ] La configuración, documentación y commit identifican el cambio.",
    ):
        add_bullet(document, text)

    add_callout(
        document,
        "Criterio de continuidad",
        "Si el candidato no supera el baseline de forma consistente, manténgalo "
        "como experimento y conserve el modelo operacional. Agregar más datos "
        "debe responder a errores observados y no solo aumentar el número bruto "
        "de imágenes.",
        fill=GREEN_FILL,
    )


def build_document() -> Document:
    document = Document()
    configure_styles(document)
    configure_sections(document)
    add_cover(document)
    add_orientation(document)
    add_prerequisites(document)
    add_data_workflow(document)
    add_training_workflow(document)
    add_operation(document)
    add_commands(document)
    add_outputs_and_troubleshooting(document)
    add_final_checklists(document)
    return document


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
